"""Load the collaborators' Excel as a benchmark and compare pipeline output against it.

Workbook layout (sheet "Quant"): a header row "AP31 (IMG_3114)" followed by three rows (band 1..3, top to
bottom) holding nine lane values: Standard 1-3 then Sample 1-6, i.e. what the web page's result table gives.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

import numpy as np
import openpyxl

_HEADER = re.compile(r"\s*(AP\d+)\s*\(IMG_(\d+)\)")


@dataclass
class BenchmarkPlate:
    plate: str
    image_no: str
    values: np.ndarray   # [band, lane]

    @property
    def image_file(self) -> str:
        return f"IMG_{self.image_no}.JPG"


def load_quant(xlsx_path, sheet: str = "Quant") -> list[BenchmarkPlate]:
    ws = openpyxl.load_workbook(xlsx_path, data_only=True)[sheet]
    plates: list[BenchmarkPlate] = []
    for r in range(2, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            m = _HEADER.match(a)
            if m:
                plates.append(BenchmarkPlate(m.group(1), m.group(2), np.full((3, 9), np.nan)))
        elif a is not None and plates:
            plates[-1].values[int(a) - 1, :] = [ws.cell(r, c).value for c in range(2, 11)]
    return plates


def compare(mine: np.ndarray, ref: np.ndarray) -> dict:
    """Agreement between two [band, lane] tables: Spearman-style rank agreement and median ratio."""
    a, b = mine.ravel(), ref.ravel()
    ra, rb = np.argsort(np.argsort(a)), np.argsort(np.argsort(b))
    strong = b > 1000
    return dict(
        rank_corr=float(np.corrcoef(ra, rb)[0, 1]),
        pearson=float(np.corrcoef(a, b)[0, 1]),
        median_ratio=float(np.median(a[strong] / b[strong])) if strong.any() else float("nan"),
    )
